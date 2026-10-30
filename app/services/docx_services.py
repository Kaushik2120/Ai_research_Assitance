
from io import BytesIO
from docx import Document


def extract_text_from_docx(file_bytes: bytes) -> list[dict]:
    document = Document(BytesIO(file_bytes))

    text_parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    # Also extract text from tables.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if cells:
                text_parts.append(" | ".join(cells))

    full_text = "\n".join(text_parts)

    return [
        {
            "page_number": 1,
            "text": full_text,
        }
    ] if full_text else []
